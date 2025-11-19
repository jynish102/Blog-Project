from django.shortcuts import render,redirect,get_object_or_404,reverse
from .form import *
from .models import *
from django.contrib.auth import logout
from openpyxl import Workbook
import datetime
from django.http import HttpResponse
# from docx import Document


def home(request):
    context = {'blogs':BlogPost.objects.all()}
    return render(request,'home.html',context)

def register_view(request):
    return render(request,'register.html')

def login_view(request):
    return render(request,'login.html')

def logout_view(request):
    logout(request)
    return redirect('/')

def add_blog(request):
    context = {'form':BlogForm}
    try:
        if request.method == 'POST':
            form = BlogForm(request.POST)
            print(request.FILES)
            image = request.FILES.get('image', '')
            title = request.POST.get('title')
            author = request.user

            if form.is_valid():
                print('Valid')
                content = form.cleaned_data['content']

            blog_obj = BlogPost.objects.create(
                author=author, title=title,
                content=content, image=image
            )
            print(blog_obj)
            return redirect('/add_blog/')
    except Exception as e:
        print(e)
         
    return render(request,'add_blog.html',context)

def blog_detail(request,slug):
    context = {}
    try:
        blog_obj = BlogPost.objects.filter(slug=slug).first()
        context['blog_obj'] = blog_obj
    except Exception as e:
        print(e)
    return render(request, 'blog_detail.html', context)

def see_blog(request):
    context = {}
    try:
        blog_objs = BlogPost.objects.filter(author=request.user)
        context['blog_objs'] = blog_objs
    except Exception as e:
        print(e)
    print(context)
    return render(request,'see_blog.html',context)

def blog_update(request,slug):
    context = {}
    try:
        
        blog_obj = BlogPost.objects.get(slug = slug)
       

        if blog_obj.author != request.user:
            return redirect('/')
        
        initial_dict = {'content':blog_obj.content}
        form = BlogForm(initial = initial_dict)
        if request.method == 'POST':
            form = BlogForm(request.POST)
            print(request.FILES)
            image = request.FILES.get('image', '')
            title = request.POST.get('title')
            author = request.user

            if form.is_valid():
                print('Valid')
                content = form.cleaned_data['content']

            blog_obj = BlogPost.objects.create(
                author=author, title=title,
                content=content, image=image
            )
           
    
        
        context['blog_obj'] = blog_obj
        context['form'] = form

    except Exception as e:
        print(e)
    return render(request,'update_blog.html',context)

def blog_delete(request,id):
    try:
        blog_obj = get_object_or_404(BlogPost, id=id)
        if blog_obj.author == request.user:
            blog_obj.delete()
            return redirect('/see_blog/')
        else:
           return redirect('/see_blog/')
    except Exception as e:
        print(e)
        return redirect('/see_blog/')

def verify(request,token):
    try:
        profile_obj = Profile.objects.filter(token = token).first()
        if profile_obj:
            profile_obj.is_verified = True
            profile_obj.save()
            return redirect('/login/')
    except Exception as e:
        print(e)  
    return redirect('/') 
               
def export_blog_excel(request):
    blogs = BlogPost.objects.all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blog Posts"

    sheet.append(["ID", "Title", "Author", "Created At", "Content"])

    for blog in blogs:
        sheet.append([blog.id, blog.title, blog.author.username, blog.created_at.strftime('%Y-%m-%d'), blog.content])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"blog_posts_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename={filename}'
    workbook.save(response)
    return response    

# def export_blog_word(request):
    blogs = BlogPost.objects.all()

    doc = Document()
    doc.add_heading("Blog Posts", level=1)

    for blog in blogs:
        doc.add_heading(blog.title, level=2)
        doc.add_paragraph(f"Author: {blog.author.username}")
        doc.add_paragraph(f"Created: {blog.created_at.strftime('%Y-%m-%d')}")
        doc.add_paragraph(blog.content)
        doc.add_paragraph("\n")

    # Response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    filename = f"blog_posts_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    response['Content-Disposition'] = f'attachment; filename={filename}'
    doc.save(response)
    return response